"""Build the public advocate directory from the two source Excel registries."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
GENERAL_FILE = ROOT / "общий список 139.xlsx"
GGUP_FILE = ROOT / "список ГГЮП 2026 г. Жетісу январь.xlsx"
OUTPUT_FILE = ROOT / "public" / "data" / "advocates-september-2026.json"


def clean_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def name_key(value: object) -> str:
    return unicodedata.normalize("NFC", clean_text(value)).casefold()


def format_kz_phone(digits: str) -> tuple[str, str] | None:
    if len(digits) != 11 or digits[0] not in {"7", "8"}:
        return None
    national = digits[1:]
    display = f"+7 ({national[:3]}) {national[3:6]}-{national[6:8]}-{national[8:]}"
    return display, f"+7{national}"


def parse_contacts(raw_value: object) -> list[dict[str, object]]:
    raw = clean_text(raw_value)
    contacts: list[dict[str, object]] = []
    for part in (clean_text(item) for item in re.split(r"[,;\n]+", raw)):
        if not part:
            continue
        digits = re.sub(r"\D", "", part)
        if len(digits) > 11 and len(digits) % 11 == 0:
            chunks = [digits[index : index + 11] for index in range(0, len(digits), 11)]
            formatted_chunks = [format_kz_phone(chunk) for chunk in chunks]
            if all(formatted_chunks):
                contacts.extend(
                    {"display": display, "href": href}
                    for display, href in formatted_chunks
                    if display and href
                )
                continue
        formatted = format_kz_phone(digits)
        if formatted:
            display, href = formatted
            contacts.append({"display": display, "href": href})
            continue

        # Preserve a valid mobile number even when the source contains an
        # unexplained fragment after it (for example, ".8707").
        leading = re.match(r"\D*([78](?:\D*\d){10})(.*)$", part)
        if leading:
            leading_digits = re.sub(r"\D", "", leading.group(1))
            mobile = format_kz_phone(leading_digits)
            if mobile:
                display, href = mobile
                contacts.append({"display": display, "href": href})
                remainder = clean_text(leading.group(2)).lstrip(". -")
                if remainder:
                    contacts.append({"display": remainder, "needsReview": True})
                continue

        # A labelled local work number is useful as supplied, but cannot be
        # converted safely to an international tel: link without an area code.
        if re.search(r"р\.?\s*т\.?​?", part, flags=re.IGNORECASE) and len(digits) == 6:
            contacts.append({"display": part})
        else:
            contacts.append({"display": part, "needsReview": True})
    return contacts


def workbook_rows(path: Path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    yield from sheet.iter_rows(min_row=4, values_only=True)


def load_general_registry() -> tuple[list[dict[str, object]], list[str]]:
    advocates: list[dict[str, object]] = []
    consultation_order: list[str] = []
    current_consultation = ""

    for source_id, region, name, contacts in workbook_rows(GENERAL_FILE):
        if isinstance(source_id, (int, float)) and clean_text(name):
            advocates.append(
                {
                    "sourceId": int(source_id),
                    "name": clean_text(name),
                    "region": clean_text(region),
                    "consultation": current_consultation,
                    "contacts": parse_contacts(contacts),
                }
            )
        elif clean_text(name):
            current_consultation = clean_text(name)
            consultation_order.append(current_consultation)

    return advocates, consultation_order


def load_ggup_names() -> list[str]:
    names: list[str] = []
    for source_id, _region, name, _contacts in workbook_rows(GGUP_FILE):
        if isinstance(source_id, (int, float)) and clean_text(name):
            names.append(clean_text(name))
    return names


def match_ggup(advocates: list[dict[str, object]], ggup_names: list[str]) -> dict[str, int]:
    canonical = {name_key(item["name"]): item for item in advocates}
    matches: dict[str, int] = {}

    for source_id, supplied_name in enumerate(ggup_names, start=1):
        supplied_key = name_key(supplied_name)
        matched_key = supplied_key if supplied_key in canonical else ""
        if not matched_key:
            scored = sorted(
                ((SequenceMatcher(None, supplied_key, candidate).ratio(), candidate) for candidate in canonical),
                reverse=True,
            )
            score, candidate = scored[0]
            if score < 0.92:
                raise ValueError(f"No safe GGUP match for {supplied_name!r}; best score={score:.3f}")
            matched_key = candidate
        if matched_key in matches:
            raise ValueError(f"Duplicate GGUP match for {supplied_name!r}")
        matches[matched_key] = source_id

    return matches


def main() -> None:
    previous = json.loads(OUTPUT_FILE.read_text(encoding="utf-8"))
    previous_ids = {name_key(item["name"]): item["id"] for item in previous["advocates"]}
    consultation_ids = {item["name"]: item["id"] for item in previous["consultations"]}

    advocates, consultation_order = load_general_registry()
    ggup_names = load_ggup_names()
    ggup_matches = match_ggup(advocates, ggup_names)

    for item in advocates:
        key = name_key(item["name"])
        item["id"] = previous_ids.get(key, "98-amantai-turlubek-zhanzakuly")
        item["ggup2026"] = key in ggup_matches
        if key in ggup_matches:
            item["ggupSourceId"] = ggup_matches[key]

    counts = Counter(str(item["consultation"]) for item in advocates)
    consultations = [
        {
            "id": consultation_ids[name],
            "name": name,
            "count": counts[name],
        }
        for name in consultation_order
    ]
    review_count = sum(
        1
        for item in advocates
        if any(contact.get("needsReview") for contact in item["contacts"])
    )

    assert len(advocates) == 139
    assert len({item["sourceId"] for item in advocates}) == 139
    assert len({item["id"] for item in advocates}) == 139
    assert len({name_key(item["name"]) for item in advocates}) == 139
    assert sum(counts.values()) == 139
    assert len(consultations) == 13
    assert sum(bool(item["ggup2026"]) for item in advocates) == 50

    directory = {
        "meta": {
            "organization": "Коллегия адвокатов области Жетісу",
            "region": "область Жетісу",
            "asOf": "2026-09-01",
            "sourceLabel": "Список адвокатов КАОЖ по состоянию на 01.09.2026",
            "sourceFile": GENERAL_FILE.name,
            "total": len(advocates),
            "consultationCount": len(consultations),
            "contactReviewCount": review_count,
            "ggup": {
                "asOf": "2026-01",
                "sourceLabel": "Список адвокатов, участвующих в ГГЮП в 2026 году",
                "sourceFile": GGUP_FILE.name,
                "total": len(ggup_names),
            },
        },
        "consultations": consultations,
        "advocates": advocates,
    }
    OUTPUT_FILE.write_text(json.dumps(directory, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"Imported {len(advocates)} advocates, {len(ggup_names)} GGUP participants; "
        f"{review_count} contact record(s) need clarification."
    )


if __name__ == "__main__":
    main()
