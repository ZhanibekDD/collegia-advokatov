#!/usr/bin/env python3
"""Import the current KAOJ member directory from the approved Excel workbook.

Usage:
  python scripts/import-advocates-xlsx.py <workbook.xlsx> [output.json]

The source workbook is intentionally kept outside the repository because it is
an administrative working document. The generated public JSON contains only
the directory fields displayed on the website.
"""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_OUTPUT = Path("public/data/advocates-september-2026.json")
EXPECTED_DATE = "2026-09-01"


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def slugify(value: str) -> str:
    transliteration = str.maketrans(
        {
            "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e",
            "ж": "zh", "з": "z", "и": "i", "й": "i", "к": "k", "л": "l", "м": "m",
            "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
            "ф": "f", "х": "h", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "shch",
            "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
            "ә": "a", "ғ": "g", "қ": "q", "ң": "n", "ө": "o", "ұ": "u", "ү": "u", "һ": "h", "і": "i",
        }
    )
    result = value.lower().translate(transliteration)
    return re.sub(r"[^a-z0-9]+", "-", result).strip("-")


def import_directory(source: Path, output: Path) -> None:
    worksheet = load_workbook(source, data_only=True, read_only=True).active
    title = clean(worksheet.cell(1, 3).value)
    if "01.09.2026" not in title:
        raise ValueError(f"Unexpected directory date in title: {title!r}")

    advocates: list[dict[str, object]] = []
    current_consultation = ""

    for row_number in range(4, worksheet.max_row + 1):
        raw_id, raw_region, raw_name = (
            worksheet.cell(row_number, column).value for column in range(1, 4)
        )

        if raw_id is None and raw_region is None and raw_name is None:
            continue

        if not isinstance(raw_id, (int, float)):
            section = clean(raw_name or raw_region or raw_id)
            if section:
                current_consultation = section
            continue

        name = clean(raw_name)
        region = clean(raw_region)
        if not name or not current_consultation:
            raise ValueError(f"Incomplete record at workbook row {row_number}")

        source_id = int(raw_id)
        advocates.append(
            {
                "id": f"{source_id}-{slugify(name)}",
                "sourceId": source_id,
                "name": name,
                "region": region,
                "consultation": current_consultation,
            }
        )

    duplicate_names = [
        name for name, count in Counter(item["name"] for item in advocates).items() if count > 1
    ]
    if duplicate_names:
        raise ValueError(f"Duplicate advocate names: {duplicate_names}")

    consultation_counts = Counter(item["consultation"] for item in advocates)
    consultations = [
        {
            "id": slugify(name),
            "name": name,
            "count": consultation_counts[name],
        }
        for name in consultation_counts
    ]

    payload = {
        "meta": {
            "organization": "Коллегия адвокатов области Жетісу",
            "region": "область Жетісу",
            "asOf": EXPECTED_DATE,
            "sourceLabel": "Общий список адвокатов КАОЖ по состоянию на 01.09.2026",
            "sourceFile": source.name,
            "total": len(advocates),
            "consultationCount": len(consultations),
        },
        "consultations": consultations,
        "advocates": advocates,
    }

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(advocates)} advocates across {len(consultations)} groups into {output}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit("Please provide the source .xlsx path")
    import_directory(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT)
